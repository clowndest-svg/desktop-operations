"""文档生成：四种壳各真写一份、再真开回来读一遍。

判据不是"没抛异常"：每种格式都用**它自己的库**读回来，看标题、小节、表头、以及
"数字是不是被写成数字"。写完就算成功，是上一轮那种"包里有文件≠能力在"的老错。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest

from jarvis.app.documents import (
    KINDS,
    DocumentError,
    DocumentService,
    capabilities,
    clean_sections,
    kind_by_id,
    missing_libraries,
    resolve_directory,
    safe_name,
    unique_path,
    write_document,
)
from jarvis.tools.builtins import document_tools
from jarvis.tools.disk_cleaner import DiskCleaner

SECTIONS: list[dict[str, Any]] = [
    {
        "heading": "一、本周",
        "paragraphs": ["把投屏控制条修好了。"],
        "bullets": ["生图接进软件", "对话里能出图"],
        "table": {
            "headers": ["事项", "状态", "耗时"],
            "rows": [["修控制条", "完成", "3"], ["写文档", "进行中", "2"]],
        },
    },
    {"heading": "二、下周", "bullets": ["联网探测"], "paragraphs": ["先量再改。"]},
]


def test_every_kind_declares_what_it_carries() -> None:
    assert [kind.id for kind in KINDS] == ["xlsx", "docx", "pptx", "pdf"]
    assert all(kind.note for kind in KINDS), "面板上那一行解释就是给用户做取舍用的"


def test_a_kind_can_be_named_by_extension_too() -> None:
    assert kind_by_id(".PDF") is not None and kind_by_id(".PDF").id == "pdf"  # type: ignore[union-attr]
    assert kind_by_id("txt") is None


class TestCleaningTheShape:
    def test_a_single_object_is_accepted_as_one_section(self) -> None:
        """模型经常给一个 dict 而不是 list —— 拒掉只会换来一次"生成失败"。"""
        assert clean_sections({"heading": "只有一节", "bullets": ["甲"]})[0].heading == "只有一节"

    def test_a_bare_string_becomes_a_paragraph(self) -> None:
        body = clean_sections("第一段\n第二段")
        assert body[0].paragraphs == ("第一段", "第二段")

    def test_a_table_may_be_a_bare_grid(self) -> None:
        row = clean_sections([{"table": [["名称", "数量"], ["苹果", "3"]]}])[0]
        assert row.headers == ("名称", "数量") and row.rows == (("苹果", "3"),)

    def test_an_empty_thing_is_refused_with_the_shape_spelled_out(self) -> None:
        with pytest.raises(DocumentError) as got:
            clean_sections([])
        assert "heading" in str(got.value) and "table" in str(got.value)


class TestNamesAndPlaces:
    def test_a_title_cannot_escape_its_directory(self) -> None:
        name = safe_name("../../windows/system32")
        assert "/" not in name and "\\" not in name and ".." not in name

    def test_an_existing_file_is_never_overwritten(self, tmp_path: Path) -> None:
        first = unique_path(tmp_path, "同一份", ".docx")
        first.write_text("旧内容", encoding="utf-8")
        second = unique_path(tmp_path, "同一份", ".docx")
        assert second != first
        assert first.read_text(encoding="utf-8") == "旧内容"

    def test_a_relative_directory_is_read_against_the_exports_dir(self, tmp_path: Path) -> None:
        exports = tmp_path / "数据/文档"
        chosen = resolve_directory("reports", exports)
        assert chosen == exports / "reports" and chosen.is_dir()

    def test_a_path_that_is_a_file_is_refused_not_replaced(self, tmp_path: Path) -> None:
        blocker = tmp_path / "挡路"
        blocker.write_text("x", encoding="utf-8")
        with pytest.raises(DocumentError) as got:
            resolve_directory(str(blocker), tmp_path / "数据/文档")
        assert "目录" in str(got.value)


class TestTheFourShells:
    """每一份都开回来读。"""

    def _path(self, tmp_path: Path, kind: str) -> Path:
        answer = write_document(kind, "小夜周报", SECTIONS, exports_dir=tmp_path)
        assert answer["ok"] is True, answer
        path = Path(str(answer["path"]))
        assert path.is_file() and path.stat().st_size > 1024, "空壳一份不算生成"
        return path

    def test_xlsx_keeps_one_sheet_per_section_and_real_numbers(self, tmp_path: Path) -> None:
        from openpyxl import load_workbook

        book = load_workbook(self._path(tmp_path, "xlsx"))
        assert book.sheetnames == ["一、本周", "二、下周"]
        cells = [
            c.value for row in book["一、本周"].iter_rows() for c in row if c.value is not None
        ]
        assert "事项" in cells and "修控制条" in cells
        assert 3 in cells and isinstance(
            cells[cells.index(3)], int
        ), "耗时被写成了文本，Excel 里就不能算"

    def test_docx_carries_headings_bullets_and_a_table(self, tmp_path: Path) -> None:
        from docx import Document

        doc = Document(str(self._path(tmp_path, "docx")))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert texts[0] == "小夜周报" and "一、本周" in texts
        assert "生图接进软件" in texts
        assert len(doc.tables) == 1
        assert [c.text for c in doc.tables[0].rows[0].cells] == ["事项", "状态", "耗时"]

    def test_pptx_opens_with_a_cover_and_one_slide_per_section(self, tmp_path: Path) -> None:
        from pptx import Presentation

        deck = Presentation(str(self._path(tmp_path, "pptx")))
        assert len(deck.slides) == 3
        first = [s.text_frame.text for s in deck.slides[0].shapes if s.has_text_frame]
        assert "小夜周报" in first[0]
        second = "\n".join(s.text_frame.text for s in deck.slides[1].shapes if s.has_text_frame)
        assert "一、本周" in second and "联网探测" not in second, "两节内容混到一页上"

    def test_pdf_is_readable_and_the_chinese_survives(self, tmp_path: Path) -> None:
        from pypdf import PdfReader

        answer = write_document("pdf", "小夜周报", SECTIONS, exports_dir=tmp_path)
        reader = PdfReader(str(answer["path"]))
        page = reader.pages[0].extract_text() or ""
        assert "小夜周报" in page and "把投屏控制条修好了" in page
        assert "进行中" in page, "表格里的中文没进来"
        if answer["font"] == "XiaoYeCJK":
            assert "生图接进软件" in page

    def test_a_machine_without_a_cjk_font_falls_back_and_says_which(
        self, tmp_path: Path, monkeypatch: Any
    ) -> None:
        monkeypatch.setattr("jarvis.app.documents.CJK_FONTS", ("Z:\\没有这个字体.ttf",))
        answer = write_document("pdf", "没字体的机器", SECTIONS, exports_dir=tmp_path)
        assert answer["font"] == "STSong-Light"
        assert Path(str(answer["path"])).is_file()


class TestTheToolLayer:
    def _handlers(self, service: DocumentService) -> dict[str, Any]:
        return {spec.name: handler for spec, handler in document_tools.build(service)}

    def test_no_service_means_no_tools_advertised(self) -> None:
        assert document_tools.build(None) == []

    def test_a_missing_library_silences_the_tools(self, tmp_path: Path, monkeypatch: Any) -> None:
        service = DocumentService(tmp_path)
        monkeypatch.setattr(service, "available", lambda: False)
        assert document_tools.build(service) == []

    def test_the_answer_names_the_path_because_the_page_reads_it_back(self, tmp_path: Path) -> None:
        handlers = self._handlers(DocumentService(tmp_path))
        words = handlers["write_document"](
            {"kind": "docx", "title": "周报", "sections": [{"heading": "一", "bullets": ["甲"]}]}
        )
        tail = words.split("存在：")[1]
        for stop in ("（", "。", "，"):
            tail = tail.split(stop)[0]
        assert tail.endswith(".docx"), tail
        assert Path(tail).is_file(), "回话里的路径当不得真"

    def test_a_refusal_is_a_sentence_not_a_traceback(self, tmp_path: Path) -> None:
        handlers = self._handlers(DocumentService(tmp_path))
        words = handlers["write_document"]({"kind": "txt", "title": "x", "sections": SECTIONS})
        assert "没有这种文档" in words and "Traceback" not in words
        assert "不能空着" in handlers["write_document"](
            {"kind": "docx", "title": "x", "sections": []}
        )
        assert "要说 kind" in handlers["write_document"]({"title": "x", "sections": SECTIONS})

    def test_the_kinds_tool_lists_all_four(self, tmp_path: Path) -> None:
        words = self._handlers(DocumentService(tmp_path))["document_kinds"]({})
        for kind in ("xlsx", "docx", "pptx", "pdf"):
            assert kind in words

    def test_capabilities_says_what_is_missing(self) -> None:
        assert missing_libraries() == () or all(
            name in capabilities() for name in missing_libraries()
        )


class _AsksForADocument:
    """只假模型那一下：第一轮调 write_document，第二轮照着工具的回话作答。"""

    provider_name = "fake"
    model = "fake"

    def __init__(self) -> None:
        self.round = 0

    def complete(self, messages: Any, *, options: Any = None) -> Any:
        from jarvis.llm.types import ChatResponse, ToolCall

        del options
        self.round += 1
        if self.round == 1:
            return ChatResponse(
                content="",
                model="fake",
                tool_calls=(
                    ToolCall(
                        id="c1",
                        name="write_document",
                        arguments=(
                            '{"kind": "docx", "title": "周报", '
                            '"sections": [{"heading": "一、本周", "bullets": ["修好了控制条"]}]}'
                        ),
                    ),
                ),
            )
        return ChatResponse(content=messages[-1].content, model="fake")

    def stream(self, messages: Any, *, options: Any = None) -> Any:
        del messages, options
        raise AssertionError("这条测试走的是 complete")


class TestThroughTheChatDoor:
    """对话那一句 → 盘上那一份文件 → 页面拿得到的那个路径。"""

    def test_a_generated_document_reaches_the_reply_as_an_artifact(self, tmp_path: Path) -> None:
        from jarvis.app.chat_service import ChatService
        from jarvis.config.schema import ToolsSection
        from jarvis.tools.monitor import SystemMonitor
        from jarvis.tools.service import ToolService

        section = ToolsSection.from_mapping(
            {
                "enabled": True,
                "confirm_dangerous": True,
                "allow_write": False,
                "allow_shell": False,
                "file_roots": [],
                "max_result_chars": 8000,
            }
        )
        exports = tmp_path / "数据" / "文档"
        service = ToolService(
            lambda: section,
            monitor_factory=lambda: SystemMonitor(top_processes=0),
            cleaner_factory=lambda: DiskCleaner(tmp_path / "audit.jsonl"),
            documents=DocumentService(exports),
        )
        service.start()
        try:
            assert "write_document" in [spec.name for spec in service.specs()]
            chat = ChatService(
                lambda: _AsksForADocument(), tool_provider=lambda: service, session_id="t"
            )
            chat.start()
            reply = chat.ask("帮我把本周情况整理成一份 Word")
            assert reply.error == "", reply.error
            assert "write_document" in reply.tools_used
            assert len(reply.artifacts) == 1, reply.artifacts
            made = Path(reply.artifacts[0])
            assert made.is_file() and made.suffix == ".docx"
            assert made.parent == exports, "没落到「数据/文档」里"
            assert reply.to_dict()["artifacts"] == [str(made)], "过桥的那一份没带上"
        finally:
            service.stop()

    def test_a_turn_that_only_mentions_a_path_produces_no_artifact(self, tmp_path: Path) -> None:
        """read_file 读到一个写着「存在：…」的文件，不该凭空多出一张图。"""
        from jarvis.app.chat_service import _artifacts_of

        fake = tmp_path / "x.png"
        fake.write_bytes(b"x")
        assert _artifacts_of("generate_image", f"画好了，存在：{fake}（512x512）") == (str(fake),)
        assert _artifacts_of("read_file", f"文件内容：存在：{fake}") == ()
        assert _artifacts_of("generate_image", "这一张没画成：本机没有答话的 ComfyUI") == ()
        assert _artifacts_of("generate_image", "画好了，存在：Z:\\没有这个文件.png（512）") == ()
